"""Buat laporan Word yang dapat diedit dan diimpor ke Google Docs."""

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / 'tmp/docx'
TMP.mkdir(parents=True, exist_ok=True)
SHOTS = ROOT / 'docs/screenshots'
IDENTITY = json.loads((ROOT / 'docs/identitas.json').read_text())
OUTPUT = TMP / 'laporan_pertemuan_3_unsanitized.docx'
LEFT = 59.025
WIDTH = 493.95
student_class = IDENTITY.get('kelas', '')
pdf_path = ROOT / 'output/pdf/laporan_pertemuan_3_152023088.pdf'
if not student_class and pdf_path.exists():
    student_class = str((PdfReader(pdf_path).get_fields() or {}).get('kelas_mahasiswa', {}).get('/V', ''))

# Ekstrak logo aslinya; seluruh teks laporan dibuat sebagai teks Word.
cover = PdfReader(ROOT / 'docs/report_assets/cover_template.pdf')
logo_image = cover.pages[0].images[0]
logo = TMP / logo_image.name
logo.write_bytes(logo_image.data)

doc = Document()
doc.core_properties.title = 'Laporan Tugas 3 Praktikum Flutter'
doc.core_properties.author = IDENTITY['nama']
doc.core_properties.subject = 'Aplikasi E-Money Scaffold dan Navigator.push()'
doc.core_properties.comments = ''
settings = doc.settings.element
update_fields = OxmlElement('w:updateFields')
update_fields.set(qn('w:val'), 'true')
settings.append(update_fields)


def style_font(style, name='Calibri', size=11, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attribute in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn(f'w:{attribute}'), name)
    for attribute in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn(f'w:{attribute}'), None)


for name, size, bold, italic in [
    ('Normal', 11, False, False),
    ('Title', 18, True, False),
    ('Heading 1', 16, True, False),
    ('Heading 2', 14, True, False),
    ('Caption', 10, False, True),
    ('Footer', 9, False, False),
]:
    style = doc.styles[name]
    style_font(style, size=size, bold=bold, italic=italic)
    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(15.25)

doc.styles['Title'].paragraph_format.line_spacing = 1.0
doc.styles['Heading 1'].paragraph_format.line_spacing = Pt(20)
doc.styles['Heading 2'].paragraph_format.line_spacing = Pt(18)
doc.styles['Caption'].paragraph_format.line_spacing = Pt(12.5)
doc.styles['Footer'].paragraph_format.line_spacing = Pt(15.25)


def configure_section(section, top):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Pt(LEFT)
    section.right_margin = Pt(LEFT)
    section.top_margin = Pt(top)
    section.bottom_margin = Pt(55)
    section.header_distance = Pt(27)
    section.footer_distance = Pt(34)


configure_section(doc.sections[0], 75)
footer = doc.sections[0].footer.paragraphs[0]
footer.style = doc.styles['Footer']
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Halaman ')
run = footer.add_run()
run.font.size = Pt(11)
begin = OxmlElement('w:fldChar')
begin.set(qn('w:fldCharType'), 'begin')
instruction = OxmlElement('w:instrText')
instruction.set(qn('xml:space'), 'preserve')
instruction.text = ' PAGE '
separate = OxmlElement('w:fldChar')
separate.set(qn('w:fldCharType'), 'separate')
value = OxmlElement('w:t')
value.text = '1'
end = OxmlElement('w:fldChar')
end.set(qn('w:fldCharType'), 'end')
for element in [begin, instruction, separate, value, end]:
    run._r.append(element)


def paragraph(value='', *, style='Normal', before=0, after=0, align=None, keep=False, container=doc):
    p = container.add_paragraph(value, style)
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = keep
    fmt.widow_control = True
    if align is not None:
        p.alignment = align
    return p


def gap(points, container=doc):
    p = paragraph(container=container)
    # Google Docs mengimpor tinggi baris absolut sebagai persentase.
    # Gunakan baris normal dan spasi paragraf agar jarak stabil saat diimpor.
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(max(0, points - 13.2))
    p.paragraph_format.widow_control = False
    p.add_run().font.size = Pt(1)
    return p


def heading(value, *, level=1, before=0, after=9):
    return paragraph(value, style=f'Heading {level}', before=before, after=after, keep=True)


def caption(value, *, container=doc, before=9, after=0):
    return paragraph(value, style='Caption', before=before, after=after,
                     align=WD_ALIGN_PARAGRAPH.CENTER, container=container)


def picture(filename, *, height=206.01, container=doc, first=None):
    p = first if first is not None else paragraph(container=container)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.font.size = Pt(1)
    shape = run.add_picture(str(SHOTS / filename), height=Pt(height))
    shape._inline.docPr.set('descr', {
        '01_menu_page.png': 'MenuPage lengkap dengan nama mahasiswa, NIM, tiga baris menu, dan navigasi bawah.',
        '02_drawer.png': 'Drawer E-Money berisi menu Beranda, Profil, dan praktikum sebelumnya.',
        '03_transfer.png': 'Halaman Transfer dengan AppBar dan tombol kembali.',
        '04_top_up.png': 'Halaman Top Up dengan AppBar dan tombol kembali.',
        '06_pulsa.png': 'Halaman Pulsa sebagai contoh navigasi menu baris kedua.',
    }[filename])
    return p


def cell_margins(cell, top=0, left=0, bottom=0, right=0):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side, points in [('top', top), ('left', left), ('bottom', bottom), ('right', right)]:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(round(points * 20)))
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    tcpr.append(margins)


def table_borders(table, visible):
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single' if visible else 'nil')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), 'D9D9D9')
        borders.append(element)
    table._tbl.tblPr.append(borders)


# Cover mengikuti logo, tabel, dan tipografi PDF contoh.
paragraph('Laporan Tugas 3 Praktikum Flutter', style='Title', align=WD_ALIGN_PARAGRAPH.CENTER)
gap(73.74)
p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.line_spacing = 1.0
p.add_run().font.size = Pt(1)
logo_run = p.add_run()
logo_run.font.size = Pt(1)
shape = logo_run.add_picture(str(logo), width=Pt(113), height=Pt(113))
shape._inline.docPr.set('descr', 'Logo Institut Teknologi Nasional Bandung')
gap(44.8)

identity_table = doc.add_table(rows=2, cols=2)
identity_table.alignment = WD_TABLE_ALIGNMENT.CENTER
identity_table.autofit = False
identity_table.columns[0].width = Pt(58.15)
identity_table.columns[1].width = Pt(156.45)
table_borders(identity_table, True)
for row_index, (label, text) in enumerate([('Nama', IDENTITY['nama']), ('NIM', IDENTITY['nim'])]):
    row = identity_table.rows[row_index]
    row.height = Pt(33.635)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for col_index, value in enumerate([label, text]):
        cell = row.cells[col_index]
        cell.width = Pt([58.15, 156.45][col_index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_margins(cell, top=0, bottom=0, left=7, right=4)
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = Pt(15.25)
        r = p.add_run(value)
        r.bold = col_index == 0
        if row_index == 1:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F3F6FA')
            cell._tc.get_or_add_tcPr().append(shade)

gap(205.0)
for value in [
    'PROGRAM STUDI INFORMATIKA',
    'FAKULTAS TEKNOLOGI INDUSTRI',
    'INSTITUT TEKNOLOGI NASIONAL',
    'BANDUNG',
    '2026',
]:
    p = paragraph(value, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(7.4)
    r = p.runs[0]
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

section = doc.add_section(WD_SECTION_START.NEW_PAGE)
configure_section(section, 55)
# Jangan biarkan paragraf penanda section menambah halaman kosong.
section_marker = doc.paragraphs[-1]
section_marker.paragraph_format.line_spacing = Pt(1)
section_marker.paragraph_format.space_before = Pt(0)
section_marker.paragraph_format.space_after = Pt(0)
section_marker.add_run().font.size = Pt(1)

# Halaman 2. Diagram pohon dibuat sebagai teks sehingga juga dapat diedit.
heading('Bukti Tampilan Aplikasi', after=5.31)
picture('01_menu_page.png')
caption('Gambar 1 Tampilan MenuPage', after=10)
paragraph('Pengujian interaksi dilakukan dengan menekan menu E-Money. Setiap menu membuka halaman tujuan berbeda menggunakan Navigator.push(). AppBar menampilkan nama dan NIM mahasiswa.', after=12)
heading('Widget Tree', level=2, after=9)
paragraph('Widget tree berikut menunjukkan susunan antarmuka dari widget akar sampai komponen pada setiap menu. Struktur ini menempatkan AppBar, Drawer, body, tombol aksi, dan navigasi bawah di dalam Scaffold.', after=12)
tree = '''MaterialApp
+-- MenuPage (StatefulWidget)
    +-- Scaffold
        +-- AppBar -> Column -> Judul, Nama, NIM
        +-- Drawer -> ListView -> DrawerHeader, ListTile
        +-- body -> SafeArea -> Center -> ConstrainedBox
        |   +-- SingleChildScrollView -> Column
        |       +-- 3 Row -> 9 Expanded
        |           +-- Material -> InkWell -> Container -> Column
        |               +-- Icon, SizedBox, Text
        +-- FloatingActionButton -> Navigator.push(TopUpPage)
        +-- BottomNavigationBar -> Beranda, Riwayat, Profil'''
p = paragraph(tree, after=9)
p.paragraph_format.line_spacing = Pt(16.5)
p.paragraph_format.left_indent = Pt(73)
p.paragraph_format.keep_together = True
p.runs[0].font.name = 'Courier New'
p.runs[0].font.size = Pt(8.5)
caption('Gambar 2 Widget tree', before=0)

# Halaman 3. Tabel tanpa garis hanya mempertahankan susunan screenshot contoh.
doc.add_page_break()
heading('Bukti Navigasi Antarhalaman', after=5.31)
shots_table = doc.add_table(rows=2, cols=2)
shots_table.alignment = WD_TABLE_ALIGNMENT.CENTER
shots_table.autofit = False
table_borders(shots_table, False)
for column in shots_table.columns:
    column.width = Pt(WIDTH / 2)
for row_index, items in enumerate([
    [('02_drawer.png', 'Gambar 3 Drawer'), ('03_transfer.png', 'Gambar 4 Halaman Transfer')],
    [('04_top_up.png', 'Gambar 5 Halaman Top Up'), ('06_pulsa.png', 'Gambar 6 Halaman Pulsa')],
]):
    for col_index, (filename, label) in enumerate(items):
        cell = shots_table.cell(row_index, col_index)
        cell.width = Pt(WIDTH / 2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cell_margins(cell)
        picture(filename, container=cell, first=cell.paragraphs[0])
        caption(label, container=cell, after=43.0 if row_index == 0 else 0)
    cant_split = OxmlElement('w:cantSplit')
    shots_table.rows[row_index]._tr.get_or_add_trPr().append(cant_split)
paragraph('Drawer menyediakan menu Beranda dan Profil. Menu Transfer, Top Up, dan Pulsa membuka halaman yang berbeda. Setiap halaman memiliki AppBar, ikon, nama halaman, serta tombol kembali ke Beranda.', before=17)
paragraph('Screenshot diambil dari aplikasi yang berjalan di Chrome pada ukuran layar 390 x 844. Seluruh menu pada baris pertama, kedua, dan ketiga dapat diklik dan kembali ke MenuPage.', before=26)

# Halaman 4. Potongan kode disimpan sebagai teks biasa, bukan screenshot kode.
doc.add_page_break()
heading('Implementasi Navigasi', after=10)
p = paragraph('lib/pages/menu_page.dart | Menu Transfer | NIM 152023088', after=12)
p.runs[0].font.size = Pt(9)
code = (SHOTS / 'navigator_push.txt').read_text().strip()
p = paragraph(code, after=0)
p.paragraph_format.left_indent = Pt(21)
p.paragraph_format.line_spacing = Pt(16)
p.paragraph_format.keep_together = True
p.runs[0].font.name = 'Courier New'
p.runs[0].font.size = Pt(10.5)
caption('Gambar 7 Potongan kode Navigator.push()', before=10)
heading('Cara Kerja Navigasi', level=2, before=18, after=9)
paragraph('Ketukan pada InkWell menjalankan onTap. Navigator.push() menambahkan MaterialPageRoute ke tumpukan halaman, lalu builder membuat TransferPage. MenuPage tetap tersimpan di bawahnya. Navigator.pop() atau tombol kembali menutup halaman tujuan dan menampilkan MenuPage lagi.')
heading('Fungsi Komponen Scaffold', level=2, before=22, after=9)
for value in [
    'Scaffold: menyediakan kerangka utama halaman aplikasi Material.',
    'AppBar: menampilkan judul E-Money, nama mahasiswa, dan NIM.',
    'body: menampilkan saldo demo serta tiga baris menu E-Money.',
    'Drawer: menyediakan navigasi samping ke Beranda dan Profil.',
    'SafeArea: menjaga konten dari notch, status bar, dan area navigasi sistem.',
    'FloatingActionButton: membuka Top Up sebagai aksi tambah saldo.',
    'BottomNavigationBar: menyediakan akses Beranda, Riwayat, dan Profil.',
]:
    paragraph(value)
heading('Hasil Pengujian', level=2, before=20, after=9)
paragraph('flutter analyze: No issues found! Tujuh pengujian widget lulus, termasuk layout layar kecil. Empat belas alur navigasi berhasil diperiksa langsung di Chrome.')
paragraph(f'Kelas: {student_class or "________________________"}', before=12)

doc.save(OUTPUT)
reopened = Document(OUTPUT)
text = '\n'.join(p.text for p in reopened.paragraphs)
for required in ['Laporan Tugas 3 Praktikum Flutter', 'Widget Tree', 'Navigator.push<void>', 'Hasil Pengujian', 'Kelas:']:
    assert required in text, required
assert len(reopened.inline_shapes) == 6
assert len(reopened.tables) == 2
assert IDENTITY['nama'] in reopened.tables[0].cell(0, 1).text
assert reopened.settings.element.find(qn('w:documentProtection')) is None
print(f'DOCX dibuat dengan teks, tabel, widget tree, dan kode yang dapat diedit: {OUTPUT}')
